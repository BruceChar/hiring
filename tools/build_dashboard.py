#!/usr/bin/env python3
"""Build the BOSS application dashboard workbook from the pipeline's JSON state.

Sheets: 投递看板(pipeline) / 状态统计(dashboard+chart) / 简历校验 / 说明
"""
import json
import os
from collections import Counter

from openpyxl import Workbook
from openpyxl.chart import BarChart, Reference
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))  # 项目根（自适应）
OUT = os.path.join(ROOT, "data", "exports", "BOSS投递看板.xlsx")

STATUSES = ["未沟通", "已沟通", "简历已发", "已回复", "约面试", "已面试", "已offer", "已拒绝"]
STATUS_FILL = {
    "未沟通": "FFF2CC", "已沟通": "DDEBF7", "简历已发": "D9E1F2", "已回复": "E2EFDA",
    "约面试": "C6E0B4", "已面试": "A9D08E", "已offer": "70AD47", "已拒绝": "F8CBAD",
}
HDR = PatternFill("solid", fgColor="16697A")
HDR_FONT = Font(color="FFFFFF", bold=True, size=11)
THIN = Border(*[Side(style="thin", color="D0D0D0")] * 4)


def load(name, default):
    """按 data/ -> data/raw/ 的顺序查找（历史产物在 raw/）。"""
    for base in (os.path.join(ROOT, "data"), os.path.join(ROOT, "data", "raw")):
        p = os.path.join(base, name)
        if os.path.exists(p):
            try:
                return json.load(open(p, encoding="utf-8"))
            except Exception:
                return default
    return default


applied = load("final_jds.json", [])
records = load("投递记录.json", [])
activity = {a["id"]: a for a in load("activity.json", [])}
validation = {v["company"]: v for v in load("validation.json", [])}
review = {r["company"]: r for r in load("resume_review.json", {}).get("reviews", [])}
replies = {r["company"]: r for r in load("replies.json", [])}
shortlist = load("shortlist.json", [])

# recruiter names from the chat list snapshot
chats = load("chats_after.json", {}).get("chats", [])
def recruiter_for(company):
    import re as _re
    for c in chats:
        line = " ".join(c)
        if company in line:
            i = line.find(company)
            prefix = line[:i].strip()
            prefix = _re.sub(r"^(\d{1,2}:\d{2}|昨天|今天|\d{2}月\d{2}日)\s*", "", prefix).strip()
            return prefix
    return ""

from datetime import datetime, timedelta, timezone
CST = timezone(timedelta(hours=8))
def local_time(iso):
    if not iso:
        return ""
    try:
        return datetime.fromisoformat(iso.replace("Z", "+00:00")).astimezone(CST).strftime("%Y-%m-%d %H:%M")
    except Exception:
        return iso[:16].replace("T", " ")

# ---- assemble rows -----------------------------------------------------------
rows = []
applied_ids = set()
for i, t in enumerate(applied, 1):
    applied_ids.add(t["id"])
    act = activity.get(t["id"], {})
    rec = next((r for r in records if r.get("company") == t["company"]), {})
    val = validation.get(t["company"], {})
    rev = review.get(t["company"], {})
    sent_img = rec.get("resume") == "image_sent"
    rep = replies.get(t["company"], {})
    if rep.get("replied") == "yes":
        status = "已回复"
    elif sent_img:
        status = "简历已发"
    elif rec.get("greet") in ("sent", "uncertain", "already_contacted"):
        status = "已沟通"
    else:
        status = "未沟通"
    rows.append({
        "来源": "本批投递", "公司": t["company"], "岗位": t["position"], "城市": t["city"],
        "薪资": act.get("salary") or t.get("salary", ""),
        "招聘者": recruiter_for(t["company"]), "活跃度": act.get("activity") or "未标注",
        "状态": status,
        "打招呼时间": local_time(rec.get("at")),
        "简历发送": "图片已发" if sent_img else "未发送",
        "简历校验": "通过" if val.get("passed") and rev.get("verdict", "pass") != "fail" else ("未通过" if (val and not val.get("passed")) or rev.get("verdict") == "fail" else "待复核"),
        "岗位链接": f"https://www.zhipin.com/job_detail/{t['id']}.html",
        "备注": (("HR已回复；" if replies.get(t["company"], {}).get("replied") == "yes" else "") + ("HR索要附件简历；" if replies.get(t["company"], {}).get("askedResume") else "") + rev.get("notes", ""))[:120],
    })
for s in shortlist:
    if s["id"] in applied_ids:
        continue
    rows.append({
        "来源": "候选池", "公司": s["company"], "岗位": s["name"], "城市": s["city"],
        "薪资": s["salary"], "招聘者": "", "活跃度": "待采集", "状态": "未沟通",
        "打招呼时间": "", "简历发送": "未发送", "简历校验": "未生成",
        "岗位链接": f"https://www.zhipin.com/job_detail/{s['id']}.html", "备注": f"{s.get('area','')} {s.get('exp','')}",
    })

wb = Workbook()
ws = wb.active
ws.title = "投递看板"
headers = list(rows[0].keys())
ws.append(headers)
for c in range(1, len(headers) + 1):
    cell = ws.cell(1, c)
    cell.fill, cell.font = HDR, HDR_FONT
    cell.alignment = Alignment(horizontal="center", vertical="center")
for r in rows:
    ws.append([r[h] for h in headers])
for ri in range(2, len(rows) + 2):
    st = ws.cell(ri, headers.index("状态") + 1).value
    ws.cell(ri, headers.index("状态") + 1).fill = PatternFill("solid", fgColor=STATUS_FILL.get(st, "FFFFFF"))
    for ci in range(1, len(headers) + 1):
        ws.cell(ri, ci).border = THIN
        ws.cell(ri, ci).alignment = Alignment(vertical="center", wrap_text=(headers[ci-1] == "备注"))
widths = {"来源": 10, "公司": 24, "岗位": 32, "城市": 8, "薪资": 14, "招聘者": 12, "活跃度": 12,
          "状态": 11, "打招呼时间": 17, "简历发送": 11, "简历校验": 11, "岗位链接": 46, "备注": 30}
for i, h in enumerate(headers, 1):
    ws.column_dimensions[get_column_letter(i)].width = widths.get(h, 14)
ws.freeze_panes = "A2"
ws.auto_filter.ref = f"A1:{get_column_letter(len(headers))}{len(rows)+1}"

# ---- dashboard ---------------------------------------------------------------
ds = wb.create_sheet("状态统计")
ds["A1"] = "BOSS直聘投递看板"; ds["A1"].font = Font(size=16, bold=True, color="16697A")
ds["A2"] = f"统计时间：{__import__('datetime').datetime.now():%Y-%m-%d %H:%M}　|　岗位总数：{len(rows)}　|　已投递：{sum(1 for r in rows if r['来源']=='本批投递')}"
ds["A3"] = "状态分布"
ds.append([])
ds.append(["状态", "数量", "占比"])
cnt = Counter(r["状态"] for r in rows)
total = len(rows)
for st in STATUSES:
    n = cnt.get(st, 0)
    ds.append([st, n, f"{n/total:.0%}" if total else "0%"])
start = ds.max_row - len(STATUSES) + 1
for i, st in enumerate(STATUSES):
    ds.cell(start + i, 1).fill = PatternFill("solid", fgColor=STATUS_FILL[st])
for row in ds.iter_rows(min_row=start - 1, max_row=ds.max_row, min_col=1, max_col=3):
    for c in row:
        c.border = THIN
        if c.row == start - 1:
            c.fill, c.font = HDR, HDR_FONT

chart = BarChart()
chart.title = "各状态岗位数量"
chart.y_axis.title = "数量"
data = Reference(ds, min_col=2, min_row=start - 1, max_row=ds.max_row)
cats = Reference(ds, min_col=1, min_row=start, max_row=ds.max_row)
chart.add_data(data, titles_from_data=True)
chart.set_categories(cats)
chart.height, chart.width = 9, 16
ds.add_chart(chart, "E3")

ds.append([])
ds.append(["活跃度分布", "数量"])
act_cnt = Counter(r["活跃度"] or "未标注" for r in rows if r["来源"] == "本批投递")
for k, v in act_cnt.most_common():
    ds.append([k, v])
ds.append([])
ds.append(["转化指标", "数值"])
ds.append(["打招呼成功率", f"{sum(1 for r in rows if r['状态'] not in ('未沟通',))/max(1,sum(1 for r in rows if r['来源']=='本批投递')):.0%}"])
ds.append(["简历发送成功率", f"{sum(1 for r in rows if r['简历发送']=='图片已发')/max(1,sum(1 for r in rows if r['来源']=='本批投递')):.0%}"])
ds.append(["简历校验通过率", f"{sum(1 for r in rows if r['简历校验']=='通过')/max(1,sum(1 for r in rows if r['来源']=='本批投递')):.0%}"])
ds.column_dimensions["A"].width = 20
ds.column_dimensions["B"].width = 12
ds.column_dimensions["C"].width = 10

# ---- validation sheet --------------------------------------------------------
vs = wb.create_sheet("简历校验")
vs.append(["公司", "确定性校验", "失败项", "独立Agent复核", "复核问题", "点评"])
for c in range(1, 7):
    vs.cell(1, c).fill, vs.cell(1, c).font = HDR, HDR_FONT
for t in applied:
    v = validation.get(t["company"], {})
    r = review.get(t["company"], {})
    vs.append([t["company"], "通过" if v.get("passed") else "未通过", "; ".join(v.get("errors", [])),
               {"pass": "通过", "fail": "未通过"}.get(r.get("verdict"), "待复核"),
               "; ".join(r.get("issues", [])), r.get("notes", "")])
for i, w in enumerate([24, 14, 40, 16, 40, 46], 1):
    vs.column_dimensions[get_column_letter(i)].width = w

# ---- notes -------------------------------------------------------------------
ns = wb.create_sheet("说明")
notes = [
    ["流程说明", ""],
    ["1. 岗位采集", "按关键词(Agent应用/AI Agent/大模型应用/LLM应用/AI应用) + 城市(杭州/上海/苏州)抓取岗位"],
    ["2. 筛选规则", "薪资区间 30-50K 重叠；杭州优先；排除已被沟通过的公司；排除销售/实习/非技术岗"],
    ["3. 活跃度优先", "优先「刚刚活跃/今日活跃」的招聘者，回复率显著更高"],
    ["4. 简历生成", "标题统一「Agent系统工程师」，正文按 JD 定制（个人简介/核心能力侧重/项目顺序）"],
    ["5. 投递前校验", "确定性校验（白底不透明、≥2000px、版块齐全、无编造）+ 独立Agent复核，未通过不投递"],
    ["6. 投递方式", "「立即沟通」发招呼；BOSS 限制对方回复前不能发附件简历，改用高清简历图即时送达"],
    ["7. 去重", "投递记录.json 记录公司/岗位/时间，重复岗位自动跳过；本表状态列为唯一事实来源"],
    ["状态口径", "未沟通→已沟通→简历已发→已回复→约面试→已面试→已offer/已拒绝"],
]
for row in notes:
    ns.append(row)
ns["A1"].font = Font(size=14, bold=True, color="16697A")
ns.column_dimensions["A"].width = 16
ns.column_dimensions["B"].width = 92
for r in ns.iter_rows(min_row=1, max_row=ns.max_row, min_col=1, max_col=2):
    for c in r:
        c.alignment = Alignment(vertical="center", wrap_text=True)

os.makedirs(os.path.dirname(OUT), exist_ok=True)
wb.save(OUT)
print("saved", OUT, "| rows:", len(rows))
print(Counter(r["状态"] for r in rows))
